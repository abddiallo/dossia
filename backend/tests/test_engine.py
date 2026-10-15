"""
Tests du moteur documentaire (aucune base de données nécessaire).
Lancer :  pytest -q tests/        (ou)   python tests/test_engine.py
Nécessite LibreOffice (soffice) et poppler-utils (pdftoppm) pour les tests de rendu.
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

from app.engine import docxutil as X  # noqa: E402
from app.engine.adapter import InventionError, analyze_posting, build, validate_rephrase, verify  # noqa: E402
from app.engine.analyzer import analyze  # noqa: E402
from app.engine.assistant import interpret  # noqa: E402
from app.engine.classify import classify  # noqa: E402
from app.engine.cv_extract import apply_diff, diff_profiles, extract_profile, extract_text, parse_period  # noqa: E402
from app.engine.filler import fill  # noqa: E402
from app.engine.library import blank_copy  # noqa: E402
from tests.fixtures.make_fixtures import build_all  # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="dossia-tests-"))
FX = build_all(TMP / "fx")
HAS_SOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))


def _profile():
    return extract_profile(FX["cv"])[0]


def _body_text(path) -> str:
    return "\n".join(X.element_text(el) for el in X.body_elements(Document(str(path))))


# ───────────── Analyse des modèles ─────────────

def test_nexo_filled_example_recognised_without_questions():
    s = analyze(FX["nexo"])
    assert s["kind"] == "filled_example" and s["status"] == "ready" and s["recognition_score"] == 1.0
    m = next(z for z in s["zones"] if z["field"] == "missions")
    assert m["kind"] == "block" and m["extra"]["detected"] == 3
    assert m["extra"]["roles"][:3] == ["header", "poste", "contexte"]
    assert s["constraints"]["anonymize"] is True


def test_helix_missions_as_table_rows():
    s = analyze(FX["helix"])
    m = next(z for z in s["zones"] if z["field"] == "missions")
    assert m["kind"] == "row_block" and m["loc"]["other_rows"] == [2, 3]
    assert next(z for z in s["zones"] if z["field"] == "nom_affiche")["sample"] == "P. L."


def test_alpha_blank_template():
    s = analyze(FX["alpha"])
    assert s["kind"] == "blank" and s["status"] == "ready"
    assert {z["field"] for z in s["zones"]} >= {"nom_affiche", "titre", "resume", "missions", "formations"}


def test_orbis_complex_then_answers_make_it_ready():
    s = analyze(FX["orbis"])
    assert s["status"] == "complex" and "zone de texte flottante" in s["complex_notes"]
    pend = {z["key"] for z in s["zones"] if z["status"] == "pending"}
    assert {"sec:outils", "sec:centres d'interet", "tx:0:0", "tx:0:1"} <= pend
    s2 = analyze(FX["orbis"], {"sec:outils": "competences", "aligned:outils": "accept", "sec:centres d'interet": "remove",
                               "tx:0:0": "disponibilite", "tx:0:1": "mobilite"})
    assert s2["status"] == "ready"


# ───────────── Remplissage fidèle ─────────────

def test_fill_nexo_preserves_formatting_and_removes_example_data():
    p = _profile()
    s = analyze(FX["nexo"])
    out = TMP / "nexo_out.docx"
    built = build(p, s, {}, None)
    fill(FX["nexo"], s, built["view"], out)
    text = _body_text(out)
    assert "J. M." not in text and "Banque de détail" not in text and "Java, Kotlin, SQL" not in text
    assert "A. D." in text and "Grand compte aéronautique · 01/2023 – 06/2024" in text
    doc = Document(str(out))
    header = next(p for p in X.iter_all_paragraphs(doc) if "Grand compte aéronautique" in X.para_text(p))
    rs = X.runs(header)
    assert X.run_is_bold(rs[0]) and not X.run_is_bold(rs[-1])  # client en gras, dates non


def test_fill_all_templates_without_leftover_placeholders():
    p = _profile()
    for key, dec in (("helix", None), ("alpha", None),
                     ("orbis", {"sec:outils": "competences", "aligned:outils": "accept", "sec:centres d'interet": "remove",
                                "tx:0:0": "disponibilite", "tx:0:1": "mobilite"})):
        s = analyze(FX[key], dec)
        out = TMP / f"{key}_out.docx"
        fill(FX[key], s, build(p, s, {}, None)["view"], out)
        t = _body_text(out)
        assert "[" not in t and "MM/AAAA" not in t, key
        assert "Trail" not in t


def test_blank_copy_for_shared_library_contains_no_personal_data():
    s = analyze(FX["nexo"])
    out = blank_copy(FX["nexo"], s, TMP / "blank.docx")
    t = _body_text(out)
    for leak in ("J. M.", "Architecte logiciel", "Banque de détail", "INSA", "Kafka"):
        assert leak not in t


# ───────────── Profil, traçabilité, assistant ─────────────

def test_cv_extraction_heuristic():
    p = _profile()
    assert p["identite"]["nom"] == "Diallo" and p["identite"]["annees_experience"] == 8
    assert [m["client"] for m in p["missions"]] == ["Airbus", "EDF", "DSNA"]
    assert p["missions"][0]["debut"] == "2023-01" and p["missions"][0]["secteur"] == "Aéronautique"
    assert parse_period("depuis mars 2022") == ("2022-03", None, True)


def test_every_segment_has_a_source_and_nothing_is_added():
    p = _profile()
    r = build(p, analyze(FX["nexo"]), {}, analyze_posting("IHM Qt 6 pour l'énergie"))
    assert r["stats"]["added"] == 0.0
    assert all(s["source_fact_ids"] for s in r["segments"] if s["final_text"])
    try:
        verify([{"field": "x", "source_fact_ids": [], "final_text": "Expert Kubernetes"}], {})
        raise AssertionError("aurait dû lever InventionError")
    except InventionError:
        pass


def test_rephrase_validator_rejects_inventions():
    assert not validate_rephrase("Outils de supervision en C++17.", "Outils de supervision en Rust.", {"C++17"})
    assert not validate_rephrase("Encadrement d'une équipe.", "Encadrement d'une équipe de 12 personnes.", set())
    assert validate_rephrase("Développement d'outils de supervision en C++17.",
                             "Supervision temps réel : développement d'outils en C++17.", {"C++17", "Temps réel"})


def test_review_question_when_sector_emphasis_differs():
    p = _profile()
    opts = {"rephrased": {"res": "Développeur C++/Qt spécialisé dans l'énergie, 8 ans d'expérience."}}
    r = build(p, analyze(FX["nexo"]), opts, analyze_posting("supervision pour un énergéticien"))
    assert r["question"] and "énergie" in r["question"]["text"]


def test_assistant_hides_and_refuses_to_invent():
    p = _profile()
    r = interpret("raccourcis la synthèse et enlève Python", p, {})
    assert r["options"]["shorten"]["res"] >= 1 and r["options"]["hidden"] == ["c3"]
    r2 = interpret("ajoute Rust", p, {})
    assert r2["propose_profile"] == {"type": "competence", "label": "Rust"} and not r2["changes"]


def test_profile_update_only_shows_differences():
    p = _profile()
    new = {**p, "missions": p["missions"] + [{**p["missions"][0], "id": "m9", "client": "Thales", "debut": "2025-01"}]}
    d = diff_profiles(p, new)
    assert len(d["missions"]) == 1 and not d["competences"]
    merged = apply_diff(p, d)
    assert len(merged["missions"]) == 4 and merged["missions"][-1]["id"] == "m4"


def test_classify_drop():
    assert classify("Modele_DC_Nexo.docx", extract_text(FX["nexo"]))[0] == "template"
    assert classify("CV_Diallo.docx", extract_text(FX["cv"]))[0] == "cv"
    posting = "Nous recherchons un développeur Qt. Démarrage ASAP, durée de la mission 6 mois, TJM selon profil."
    assert classify("besoin.txt", posting)[0] == "job_posting"


# ───────────── Modèles « formulaire » (structures de vrais modèles d'ESN) ─────────────

def _profile_with_levels():
    p = _profile()
    lv = {"C++17": 5, "QML": 5, "Git": 4, "Linux": 3}
    for c in p["competences"]:
        c["niveau"] = lv.get(c["label"])
    p["identite"].update(mobilite="Toulouse", disponibilite="Immédiate")
    p["certifications"] = [{"id": "ce1", "intitule": "ISTQB Foundation", "annee": "2020"}]
    return p


def test_form_template_recognised():
    s = analyze(FX["formulaire"])
    assert s["status"] == "ready", [(z["key"], z["status"]) for z in s["zones"]]
    kinds = {(z["kind"], z["field"]) for z in s["zones"]}
    assert ("labels", "nom,mobilite") in kinds and ("cell_value", "experience_annees") in kinds
    assert ("comp_table", "competences") in kinds and ("field", "trigramme") in kinds
    m = next(z for z in s["zones"] if z["field"] == "missions")
    assert m["kind"] == "block" and "dates" in m["extra"]["roles"] and "bullet" in m["extra"]["roles"]
    assert s["constraints"]["anonymize"] is False or "trigramme" in str(s["zones"])


def test_form_template_filled():
    p = _profile_with_levels()
    s = analyze(FX["formulaire"])
    out = TMP / "formulaire_out.docx"
    fill(FX["formulaire"], s, build(p, s, {}, None)["view"], out)
    t = _body_text(out)
    for ph in ("X années", "YYYY", "Lorem ipsum", "MM/YYYY", "CLIENT", "Ville - Pays", "Activité 1"):
        assert ph not in t, ph
    assert "Nom: DIALLO" in t and "Mobilité géographique: Toulouse" in t and "8 années" in t
    assert "2020 : ISTQB Foundation" in t and "Anglais (professionnel)" in t
    assert "Développeur IHM - Airbus" in t and "01/2023 – 06/2024" in t
    doc = Document(str(out))
    hdr = [p.text for p in doc.sections[0].header.paragraphs]
    assert "ADI" in hdr and "Développeur C++ / Qt / QML" in hdr
    grid = X.body_elements(doc)
    levels = [X.cell_text(c) for el in grid if el.tag == X.W_TBL for tr in X.table_rows(el) for c in X.row_cells(tr)]
    assert "3" in levels  # C++17 niveau 5/5 → 3 sur l'échelle 1-3 du modèle


def test_unknown_level_is_never_invented():
    from app.engine.formfill import _level_text
    assert _level_text(None, 3) == "" and _level_text(5, 3) == "3" and _level_text(3, 3) == "2"


def test_real_templates_if_available():
    """Modèles réels d'ESN (non fournis dans le dépôt) : DOSSIA_REAL_TEMPLATES=/chemin/vers/modeles."""
    import os
    d = os.environ.get("DOSSIA_REAL_TEMPLATES")
    if not d:
        return
    p = _profile_with_levels()
    for f in sorted(Path(d).glob("*.docx")):
        s = analyze(f)
        assert s["status"] == "ready", f.name
        out = TMP / f"real_{f.stem}.docx"
        fill(f, s, build(p, s, {}, None)["view"], out)
        t = _body_text(out)
        assert not any(x in t for x in ("XXXX", "MM/YYYY", "MM/AAAA", "Lorem ipsum")), f.name


# ───────────── Rendu complet (LibreOffice) ─────────────

def test_generate_respects_page_limit():
    if not HAS_SOFFICE:
        return
    from app.engine.generate import generate
    s = analyze(FX["nexo"])
    s["constraints"]["max_pages"] = 1
    r = generate(_profile(), FX["nexo"], s, {}, None, TMP / "gen", make_previews=False)
    assert r["pages"] <= 1 and r["adjustments"] and r["fidelity"]["level"] in ("ok", "warn")


def test_match_offer_spans_and_missing():
    from app.engine.match import match
    p = _profile()
    text = "Développement en Qt 6 et C++17 sous Linux. Connaissance de Kubernetes appréciée."
    r = match(p, text)
    labels = {s["label"]: s["have"] for s in r["skills"]}
    assert labels.get("Kubernetes") is False
    for s in r["spans"]:  # chaque position pointe sur le bon passage du texte d'origine
        assert text[s["start"]:s["end"]].strip()
    assert r["total"] == len(r["skills"]) and 0 <= r["have"] <= r["total"]


def test_profile_assistant_never_invents():
    from app.engine.profile_assistant import interpret as say
    p = _profile()
    r = say("disponible à partir du 1er décembre", p)
    assert r["profile"]["identite"]["disponibilite"] == "À partir du 1er décembre"
    r = say("bonjour", p)
    assert not r["changes"] and r["profile"] == p


def test_highlight_copy_keeps_text():
    from app.engine.generate import generate
    from app.engine.highlight import highlight_copy
    s = analyze(FX["nexo"])
    r = generate(_profile(), FX["nexo"], s, {}, None, TMP / "hl", make_previews=False)
    out = TMP / "hl" / "surligne.docx"
    n = highlight_copy(r["docx"], out, ["C++", "Qt"])
    before = "".join(X.para_text(p) for p in Document(str(r["docx"])).element.body.iter(X.qn("w:p")))
    after = "".join(X.para_text(p) for p in Document(str(out)).element.body.iter(X.qn("w:p")))
    assert n > 0 and before == after


if __name__ == "__main__":
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"✓ {name}")
            except Exception as e:  # noqa: BLE001
                fails += 1
                print(f"✗ {name} : {type(e).__name__} {e}")
    print(f"\n{'ÉCHEC' if fails else 'OK'} — {fails} échec(s)")
    sys.exit(1 if fails else 0)
