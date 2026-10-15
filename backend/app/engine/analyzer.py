"""
Analyse d'un modèle d'ESN (.docx) → « spec » JSON décrivant les zones à remplir.  [D2]

Le principe : on ne réécrit jamais le document. On repère où se trouvent les informations
(identité, synthèse, compétences, missions…) et comment elles sont mises en forme, pour que
le remplissage réutilise les paragraphes, tableaux et styles d'origine.

Gère :
  - modèles vides (placeholders [Titre], MM/AAAA…) et dossiers remplis d'un autre consultant ;
  - missions en paragraphes, une mission par tableau, ou un tableau avec une ligne par mission ;
  - compétences en tableau, en liste, ou sur une ligne ;
  - zones de texte flottantes, sections inconnues, listes alignées avec des espaces (→ questions).
"""
from __future__ import annotations

import hashlib
import re
from collections import Counter
from pathlib import Path

from docx import Document

from . import docxutil as X
from . import forms as F
from . import formscan as FS
from .text import (
    GENERIC_CLIENT_WORDS, SECTION_LABELS, YEARS_EXP_RE, has_date_range, looks_like_initials,
    looks_like_person_name, norm, placeholder_like, section_of_heading,
)

AUTO_THRESHOLD = 0.85
SPEC_VERSION = 1

ENV_RE = re.compile(r"^(environnement( technique)?|technologies|outils|stack|env\.?|langages et outils)\s*:", re.I)
CTX_RE = re.compile(r"^(contexte|projet|mission|objectif|descriptif)\s*:", re.I)

FIELD_CHOICES = [
    {"value": "competences", "label": "Compétences"},
    {"value": "missions", "label": "Expériences / missions"},
    {"value": "formations", "label": "Formations"},
    {"value": "certifications", "label": "Certifications"},
    {"value": "langues", "label": "Langues"},
    {"value": "keep", "label": "Garder tel quel"},
    {"value": "remove", "label": "Supprimer la section"},
]
TEXTBOX_CHOICES = [
    {"value": "disponibilite", "label": "Ma disponibilité"},
    {"value": "mobilite", "label": "Ma mobilité"},
    {"value": "titre", "label": "Mon titre"},
    {"value": "keep", "label": "Garder tel quel"},
    {"value": "clear", "label": "Vider"},
]


class _Ctx:
    def __init__(self, doc, decisions: dict | None = None):
        self.doc = doc
        self.decisions = decisions or {}
        self.els = X.body_elements(doc)
        self.zones: list[dict] = []
        self._n = 0
        self.claimed_el: set[int] = set()            # éléments déjà attribués à une zone
        self.claimed_p: set[tuple[int, int]] = set()  # (élément, n° de paragraphe) déjà attribués

    def keys(self) -> set[str]:
        return {z["key"] for z in self.zones}

    def _claim(self, kind: str, loc: dict) -> None:
        el = loc.get("el")
        if el is not None and "p" in loc:
            self.claimed_p.add((el, loc["p"]))
            for k in loc.get("extra_ps", []):
                self.claimed_p.add((el, k))
        elif el is not None and kind in ("table_rows", "row_block", "comp_table"):
            self.claimed_el.add(el)
        for key in ("els", "others", "extra_els"):
            self.claimed_el.update(loc.get(key, []))
        if "heading" in loc:
            self.claimed_el.add(loc["heading"])
        for e, k in loc.get("paras", []):
            self.claimed_p.add((e, k))

    def zone(self, **kw) -> dict:
        self._n += 1
        z = {"id": f"z{self._n}", "status": "auto", "question": None, "choices": None, **kw}
        z.setdefault("key", z["id"])
        z.setdefault("extra", {})
        self._claim(z["kind"], z.get("loc") or {})
        if z["key"] in self.decisions:
            z["status"], z["answer"] = "confirmed", self.decisions[z["key"]]
        elif z["confidence"] < AUTO_THRESHOLD and z["status"] == "auto":
            z["status"] = "pending"
        self.zones.append(z)
        return z

    def text(self, i: int) -> str:
        return X.element_text(self.els[i])


# ───────────────────────────── Titres de sections ─────────────────────────────

def _heading_like(ctx: _Ctx, i: int) -> tuple[bool, str]:
    el = ctx.els[i]
    if el.tag == X.W_P:
        t = X.para_text(el).strip()
        if not t or len(t) > 50 or has_date_range(t) or placeholder_like(t):
            return False, t
        st = X.style_name(el, ctx.doc).lower()
        styled = (st.startswith(("heading", "titre")) or "heading" in st) and "title" != st
        upper = t.isupper() and len(t) > 3
        return (styled or upper or X.para_is_bold(el)), t
    # bandeau = tableau d'une ligne avec un seul texte court (souvent une icône dans la cellule voisine)
    rows = X.table_rows(el)
    if len(rows) == 1:
        texts = [X.cell_text(c).strip() for c in X.row_cells(rows[0])]
        filled = [t for t in texts if t]
        if len(filled) == 1 and len(filled[0]) <= 60 and "\n" not in filled[0]:
            return True, filled[0]
    return False, ""


def _strong_heading(ctx: _Ctx, i: int, t: str) -> bool:
    """Titre inconnu : uniquement style « Titre n » ou MAJUSCULES (évite de prendre un poste pour un titre)."""
    el = ctx.els[i]
    if el.tag != X.W_P:
        return True
    st = X.style_name(el, ctx.doc).lower()
    return st.startswith(("heading", "titre")) or (t.isupper() and len(t) > 3)


def _find_sections(ctx: _Ctx) -> list[dict]:
    secs, first_known = [], None
    likes = [_heading_like(ctx, i) for i in range(len(ctx.els))]
    repeated = Counter(norm(t) for like, t in likes if like and t)
    for i in range(len(ctx.els)):
        like, t = likes[i]
        known = section_of_heading(t) if t and not placeholder_like(t) else None
        if known and (like or len(t) < 30):
            secs.append({"heading": i, "type": known, "title": t})
            first_known = i if first_known is None else first_known
        elif like and first_known is not None and t and not t.endswith(":") and len(t.split()) <= 5 \
                and _strong_heading(ctx, i, t) and repeated[norm(t)] == 1 \
                and F.mission_field(t) not in ("client", "poste", "dates") and not F.is_placeholder(t):
            secs.append({"heading": i, "type": None, "title": t})
    for k, s in enumerate(secs):
        end = secs[k + 1]["heading"] if k + 1 < len(secs) else len(ctx.els)
        s["content"] = list(range(s["heading"] + 1, end))
    return secs


# ───────────────────────────── Identité (avant la 1re section) ─────────────────────────────

def _identity(ctx: _Ctx, stop: int) -> None:
    cands = []
    for i in range(stop):
        if i in ctx.claimed_el:
            continue
        for k, p in enumerate(X.element_paragraphs(ctx.els[i])):
            t = X.para_text(p).strip()
            if not t or (i, k) in ctx.claimed_p:
                continue
            if F.is_placeholder(t) or (F.inline_labels(t) and ":" in t):
                continue
            n = norm(t)
            if re.search(r"dossier de competences|curriculum|^cv$|profil consultant|dossier technique", n):
                continue
            if X.style_name(p, ctx.doc).lower() in ("title", "titre"):
                continue  # logo / nom de l'ESN
            cands.append((i, k, t))

    def name_rank(t: str) -> int:
        n = norm(t)
        if placeholder_like(t) and re.search(r"\bnom\b|pr[ée]nom|consultant", n):
            return 4
        if looks_like_initials(t):
            return 3
        if looks_like_person_name(t) and not t.isupper():
            return 2
        if looks_like_person_name(t):
            return 1
        return 0

    used: set[tuple[int, int]] = set()
    exp = next((c for c in cands if (YEARS_EXP_RE.search(c[2]) or re.search(r"\[[^\]]*\]\s*ans", c[2])) and len(c[2]) < 60), None)
    for i, k, t in cands:
        n = norm(t)
        if "disponib" in n:
            used.add((i, k))
            ctx.zone(kind="field", field="disponibilite", loc={"el": i, "p": k}, confidence=0.95, sample=t,
                     label="Disponibilité", extra={"mode": "after_label"}, key=f"id:dispo")
        elif "mobilit" in n:
            used.add((i, k))
            ctx.zone(kind="field", field="mobilite", loc={"el": i, "p": k}, confidence=0.95, sample=t,
                     label="Mobilité", extra={"mode": "after_label"}, key=f"id:mobilite")
    if exp and "id:exp" not in ctx.keys():
        used.add(exp[:2])
        ctx.zone(kind="field", field="experience_annees", loc={"el": exp[0], "p": exp[1]}, confidence=0.93,
                 sample=exp[2], label="Années d'expérience", extra={"mode": "years"}, key="id:exp")
    ranked = [(name_rank(t), idx) for idx, (_, _, t) in enumerate(cands) if (cands[idx][0], cands[idx][1]) not in used]
    ranked = [r for r in ranked if r[0] > 0]
    if not ranked:
        return
    best = max(r[0] for r in ranked)
    name_idx = min(idx for rk, idx in ranked if rk == best)
    i, k, t = cands[name_idx]
    used.add((i, k))
    ctx.zone(kind="field", field="nom_affiche", loc={"el": i, "p": k}, confidence=0.95 if best >= 2 else 0.8, sample=t,
             label="Nom / initiales", extra={"mode": "replace", "initials": looks_like_initials(t)}, key="id:nom")
    for i2, k2, t2 in cands[name_idx + 1:]:
        if (i2, k2) in used or len(t2) >= 120 or "id:titre" in ctx.keys():
            continue
        ctx.zone(kind="field", field="titre", loc={"el": i2, "p": k2}, confidence=0.92, sample=t2,
                 label="Titre", extra={"mode": "replace"}, key="id:titre")
        break


# ───────────────────────────── Missions ─────────────────────────────

def mission_roles(paras: list, doc) -> list[str]:
    """Rôle de chaque paragraphe dans un bloc mission (version historique, voir formscan.block_roles)."""
    roles: list[str | None] = []
    for p in paras:
        t = X.para_text(p).strip()
        if not t:
            roles.append("empty")
        elif has_date_range(t):
            rest = re.sub(r"[\s·\-–—|,:/()]+", " ", _strip_dates(t)).strip()
            roles.append("header" if len(rest) >= 2 else "dates")
        elif ENV_RE.match(norm(t) + ":") or ENV_RE.match(t):
            roles.append("env")
        elif X.is_list_paragraph(p, doc):
            roles.append("bullet")
        else:
            roles.append(None)
    has_dates_only = "dates" in roles
    has_client = "header" in roles
    has_poste = has_ctx = False
    for idx, r in enumerate(roles):
        if r is not None:
            continue
        t = X.para_text(paras[idx]).strip()
        near_dates = has_dates_only and any(
            roles[j] == "dates" for j in (idx - 1, idx + 1) if 0 <= j < len(roles))
        if not has_client and near_dates and len(t) < 70:
            roles[idx], has_client = "client", True
        elif not has_poste and len(t) < 90 and not CTX_RE.match(t):
            roles[idx], has_poste = "poste", True
        elif not has_ctx:
            roles[idx], has_ctx = "contexte", True
        else:
            roles[idx] = "bullet" if idx and roles[idx - 1] == "bullet" else "other"
    return [r or "other" for r in roles]


def _strip_dates(t: str) -> str:
    from .text import DATE_RANGE_RE
    return DATE_RANGE_RE.sub(" ", t)


def _missions(ctx: _Ctx, sec: dict) -> None:
    content = [i for i in sec["content"]]
    if not content:
        return
    # Cas 1 : un tableau dont plusieurs lignes contiennent une période → une ligne par mission
    # (sauf si plusieurs tableaux datés : c'est alors un tableau par mission → cas 2)
    dated_tables = [i for i in content if ctx.els[i].tag == X.W_TBL and has_date_range(ctx.text(i))]
    for i in (content if len(dated_tables) < 2 else []):
        el = ctx.els[i]
        if el.tag != X.W_TBL:
            continue
        rows = X.table_rows(el)
        date_rows = [r for r, tr in enumerate(rows) if has_date_range("\n".join(X.cell_text(c) for c in X.row_cells(tr)))]
        if len(date_rows) >= 2 or (len(date_rows) == 1 and len(rows) >= 2 and date_rows[0] > 0 and _is_header_row(rows[0])):
            proto = rows[date_rows[0]]
            paras = [p for tc in X.row_cells(proto) for p in tc.iterchildren(X.qn("w:p"))]
            roles = FS.para_block_roles(paras, ctx.doc)
            ctx.zone(kind="row_block", field="missions", confidence=0.95 if len(date_rows) >= 2 else 0.88,
                     loc={"el": i, "proto_row": date_rows[0], "other_rows": date_rows[1:]},
                     sample=X.element_text(proto)[:200], label="Bloc mission (lignes de tableau)", key="missions",
                     extra={"roles": roles, "detected": len(date_rows)})
            return
    # Cas 2 : blocs successifs (paragraphes ou un tableau par mission)
    anchors = [i for i in content if has_date_range(ctx.text(i))]
    if not anchors:
        ctx.zone(kind="section_unknown", field=None, confidence=0.4, loc={"heading": sec["heading"], "els": content},
                 key="sec:" + norm(sec["title"]),
                 sample=" ".join(ctx.text(i) for i in content)[:200], label=f"Section « {sec['title']} »",
                 question=f"Aucune mission reconnue sous « {sec['title']} ». Que contient cette section ?",
                 choices=FIELD_CHOICES)
        return
    first = anchors[0]
    lead = [i for i in content if i < first and not X.is_empty_para(ctx.els[i])]
    pre = min(len(lead), 2)
    starts = []
    for a in anchors:
        s = a
        back = 0
        while back < pre and s - 1 >= content[0] and (not starts or s - 1 >= starts[-1] + 1) \
                and not has_date_range(ctx.text(s - 1)):
            s -= 1
            if not X.is_empty_para(ctx.els[s]):
                back += 1
        starts.append(s)
    blocks = []
    for k, s in enumerate(starts):
        end = starts[k + 1] if k + 1 < len(starts) else content[-1] + 1
        blocks.append(list(range(s, end)))
    proto = blocks[0]
    paras = [p for i in proto for p in X.element_paragraphs(ctx.els[i])]
    roles = FS.block_roles([ctx.els[i] for i in proto], ctx.doc)
    same = len(blocks) >= 2 and _role_signature(ctx, blocks[1]) == _role_signature_from(roles)
    conf = 0.96 if same else (0.88 if len(blocks) == 1 else 0.86)
    ctx.zone(kind="block", field="missions", confidence=conf,
             loc={"els": proto, "others": [i for b in blocks[1:] for i in b]},
             sample=" / ".join(X.para_text(p).strip() for p in paras if X.para_text(p).strip())[:220],
             label=f"Bloc mission ×{len(blocks)} · répétable", key="missions", extra={"roles": roles, "detected": len(blocks)})
    intro = [i for i in content if i < proto[0]]
    if intro and any(not X.is_empty_para(ctx.els[i]) for i in intro):
        ctx.zone(kind="remove", field=None, confidence=0.9, loc={"els": intro}, sample=" ".join(ctx.text(i) for i in intro)[:120],
                 label="Texte d'introduction des missions", extra={}, key="missions:intro")


def _is_header_row(tr) -> bool:
    texts = [X.cell_text(c).strip() for c in X.row_cells(tr)]
    return all(len(t) < 40 for t in texts) and not any(has_date_range(t) for t in texts)


def _role_signature(ctx: _Ctx, block: list[int]) -> tuple:
    return _role_signature_from(FS.block_roles([ctx.els[i] for i in block], ctx.doc))


def _role_signature_from(roles: list[str]) -> tuple:
    out = []
    for r in roles:
        if r in ("empty", "other", "keep", "drop"):
            continue
        if out and out[-1] == r == "bullet":
            continue
        out.append(r)
    return tuple(out)


# ───────────────────────────── Listes (compétences, formations…) ─────────────────────────────

def _items(ctx: _Ctx, content: list[int]) -> list[int]:
    return [i for i in content if ctx.els[i].tag == X.W_P and X.para_text(ctx.els[i]).strip()]


def _list_format(samples: list[str]) -> dict:
    s = " ".join(samples)
    return {
        "sep": " — " if "—" in s else (" – " if " – " in s else (" - " if " - " in s else " — ")),
        "colon": ":" in s,
        "paren_year": bool(re.search(r"\((?:19|20)\d{2}|ann[ée]e|date\)", s, re.I)),
        "rating": next((c for c in "●★■" if c in s), None),
    }


def _competences(ctx: _Ctx, sec: dict, complex_notes: list) -> None:
    content = sec["content"]
    tables = [i for i in content if ctx.els[i].tag == X.W_TBL]
    legend = " ".join(X.para_text(ctx.els[i]) for i in content if ctx.els[i].tag == X.W_P and re.search(r"\d\s*-", X.para_text(ctx.els[i])))
    form_tables = [(i, d) for i in tables if (d := FS.competence_table(ctx, i, legend))]
    if form_tables:
        for i, d in form_tables:
            labels = " / ".join(dict.fromkeys(e["label"] for e in d["rows"]))
            ctx.zone(kind="comp_table", field="competences", confidence=0.93, key=f"competences:{i}", loc={"el": i},
                     sample=labels[:200], label=f"Compétences : {labels[:60]}", extra=d)
        return
    if tables:
        i = tables[0]
        rows = X.table_rows(ctx.els[i])
        ncols = max(len(X.row_cells(r)) for r in rows)
        first = 1 if len(rows) > 1 and _looks_like_table_header(rows[0]) else 0
        if first < len(rows):
            samples = [X.element_text(r) for r in rows[first:]]
            field = "competences_par_categorie" if ncols >= 2 else "competences_flat"
            ctx.zone(kind="table_rows", field=field, confidence=0.94,
                     loc={"el": i, "first_row": first, "last_row": len(rows) - 1},
                     sample=" / ".join(samples)[:200], label="Tableau compétences", key="competences",
                     extra={"cols": ncols, **_list_format(samples)})
            return
    items = _items(ctx, content)
    if not items:
        return
    samples = [X.para_text(ctx.els[i]).strip() for i in items]
    aligned = any(re.search(r"\S {3,}\S|\t", s) for s in samples)
    if len(items) == 1 and ("," in samples[0] or placeholder_like(samples[0])) and ":" not in samples[0]:
        ctx.zone(kind="list", field="competences_inline", confidence=0.92, loc={"els": items},
                 sample=samples[0][:200], label="Compétences (sur une ligne)", key="competences", extra={"inline": True, **_list_format(samples)})
    elif aligned and ctx.decisions.get("aligned:" + norm(sec["title"])) == "keep":
        return
    elif aligned:
        ctx.zone(key="aligned:" + norm(sec["title"]), kind="list", field="competences_par_categorie", confidence=0.7, loc={"els": items},
                 sample=" / ".join(samples)[:200], label="Liste alignée avec des espaces",
                 question="Cette liste est alignée avec des espaces. On la recrée proprement (une ligne par catégorie) ?",
                 choices=[{"value": "accept", "label": "Oui, recréer proprement"}, {"value": "keep", "label": "Garder tel quel"}],
                 extra={"aligned": True, **_list_format(samples)})
        complex_notes.append("liste alignée avec des espaces")
    else:
        colon = sum(":" in s for s in samples) >= max(1, len(samples) // 2)
        ctx.zone(kind="list", field="competences_par_categorie" if colon else "competences_flat", confidence=0.92,
                 loc={"els": items}, sample=" / ".join(samples)[:200], label="Compétences (liste)", key="competences",
                 extra=_list_format(samples))


def _looks_like_table_header(tr) -> bool:
    texts = [norm(X.cell_text(c)) for c in X.row_cells(tr)]
    bold = all(X.para_is_bold(p) for c in X.row_cells(tr) for p in c.iterchildren(X.qn("w:p")) if X.para_text(p).strip())
    kw = any(re.search(r"domaine|categorie|competence|niveau|technolog|type|savoir|outil", t) for t in texts)
    return (bold or kw) and all("," not in t for t in texts)


def _simple_list(ctx: _Ctx, sec: dict, field: str) -> None:
    content = sec["content"]
    tables = [i for i in content if ctx.els[i].tag == X.W_TBL]
    tables = [i for i in tables if not F.identity_field(X.cell_text(X.row_cells(X.table_rows(ctx.els[i])[0])[0]) or "-")]
    if tables and not _items(ctx, content):
        i = tables[0]
        rows = X.table_rows(ctx.els[i])
        first = 1 if len(rows) > 1 and _looks_like_table_header(rows[0]) else 0
        last = max((r for r, tr in enumerate(rows) if X.element_text(tr).strip()), default=len(rows) - 1)
        ctx.zone(kind="table_rows", field=field, confidence=0.9, loc={"el": i, "first_row": first, "last_row": last},
                 sample=X.element_text(ctx.els[i])[:200], label=SECTION_LABELS.get(field, field), key=field,
                 extra={"cols": max(len(X.row_cells(r)) for r in rows), **F.list_format([X.element_text(r) for r in rows])})
        return
    items = _items(ctx, content)
    if items:
        samples = [X.para_text(ctx.els[i]).strip() for i in items]
        ctx.zone(kind="list", field=field, confidence=0.93, loc={"els": items}, sample=" / ".join(samples)[:200],
                 label=SECTION_LABELS.get(field, field), extra=F.list_format(samples), key=field)


def _resume(ctx: _Ctx, sec: dict) -> None:
    items = _items(ctx, sec["content"])
    if items:
        ctx.zone(kind="field", field="resume", confidence=0.95,
                 loc={"el": items[0], "p": 0, "extra_els": items[1:]},
                 sample=" ".join(X.para_text(ctx.els[i]).strip() for i in items)[:300], label="Synthèse", key="resume",
                 extra={"mode": "replace"})


# ───────────────────────────── Zones de texte flottantes ─────────────────────────────

def _textboxes(ctx: _Ctx, complex_notes: list) -> None:
    for n, tx in enumerate(X.textboxes(ctx.doc)):
        for k, p in enumerate(tx.iterchildren(X.qn("w:p"))):
            t = X.para_text(p).strip()
            if not t:
                continue
            nt = norm(t)
            key = f"tx:{n}:{k}"
            dec = ctx.decisions.get(key)
            if dec == "keep":
                continue
            guess = "disponibilite" if "disponib" in nt else ("mobilite" if "mobilit" in nt else None)
            if dec in ("disponibilite", "mobilite", "titre"):
                guess = dec
            mode = "clear" if dec == "clear" else ("after_label" if ":" in t else "replace")
            ctx.zone(key=key, kind="textbox", field=None if dec == "clear" else guess, confidence=0.75 if guess else 0.5, loc={"tx": n, "p": k},
                     sample=t, label="Zone flottante",
                     question=f"Zone flottante « {t[:40]} » : que doit-on y mettre ?",
                     choices=TEXTBOX_CHOICES, extra={"mode": mode, "suggested": guess})
        complex_notes.append("zone de texte flottante")


# ───────────────────────────── Point d'entrée ─────────────────────────────

def analyze(path: str | Path, decisions: dict | None = None) -> dict:
    """
    decisions : réponses de l'utilisateur (ou de la bibliothèque partagée [D4]), par clé de zone :
      "sec:<titre normalisé>" → competences|missions|formations|certifications|langues|keep|remove
      "tx:<n>:<k>"            → disponibilite|mobilite|titre|keep|clear
      "aligned:<titre>"       → accept|keep
    """
    doc = Document(str(path))
    ctx = _Ctx(doc, decisions)
    complex_notes: list[str] = []
    secs = _find_sections(ctx)
    first_heading = secs[0]["heading"] if secs else len(ctx.els)
    for sec in secs:
        ctx.claimed_el.add(sec["heading"])
    seen: set[str] = set()
    for sec in secs:
        t = sec["type"]
        dec = ctx.decisions.get("sec:" + norm(sec["title"]))
        if dec == "keep":
            continue
        if dec == "remove":
            ctx.zone(kind="remove", field=None, confidence=1.0, loc={"els": [sec["heading"], *sec["content"]]},
                     sample=sec["title"], label=f"Section « {sec['title']} » supprimée", key="sec:" + norm(sec["title"]))
            continue
        if dec in SECTION_LABELS:
            t = dec
        tn = norm(sec["title"])
        if "formation" in tn and "certif" in tn:
            t = "formations_certifications"
        if t and t in seen and t != "missions":
            t = None
        if t:
            seen.add(t)
        if t == "resume":
            _resume(ctx, sec)
        elif t == "competences":
            _competences(ctx, sec, complex_notes)
        elif t == "missions":
            _missions(ctx, sec)
        elif t in ("formations", "certifications", "langues", "formations_certifications"):
            _simple_list(ctx, sec, t)
        elif sec["content"]:
            ctx.zone(kind="section_unknown", field=None, confidence=0.45, key="sec:" + norm(sec["title"]),
                     loc={"heading": sec["heading"], "els": sec["content"]},
                     sample=" ".join(ctx.text(i) for i in sec["content"])[:200],
                     label=f"Section « {sec['title']} »",
                     question=f"Section « {sec['title']} » : que doit-on en faire ?", choices=FIELD_CHOICES)
    FS.form_scan(ctx)                     # formulaires, lignes libellé/valeur, en-têtes
    _identity(ctx, first_heading)         # identité « classique » sur ce qui reste
    _textboxes(ctx, complex_notes)
    # Jamais d'impasse : si aucune zone missions n'est trouvée, on pose la question.
    if not any(z["field"] == "missions" for z in ctx.zones) and ctx.decisions.get("missions:missing") != "skip":
        ctx.zone(kind="missing", field=None, confidence=0.3, loc={}, key="missions:missing",
                 label="Missions introuvables", sample="",
                 question="Je n'ai pas trouvé où placer vos missions dans ce modèle. Que faire ?",
                 choices=[{"value": "skip", "label": "Générer quand même (sans les missions)"},
                          {"value": "support", "label": "Le faire analyser par l'équipe"}])
    if any(True for _ in doc.element.body.iter(X.qn("w:drawing"))) and not X.textboxes(doc):
        pass  # images (logos) conservées telles quelles

    zones = ctx.zones
    kind = _template_kind(zones)
    for z in zones:  # suggestions par défaut selon le type de modèle
        if z["kind"] == "section_unknown":
            z["suggested"] = "remove" if kind == "filled_example" else "keep"
    scored = [z for z in zones if z["kind"] != "remove"]
    auto = [z for z in scored if z["status"] in ("auto", "confirmed")]
    score = round(len(auto) / len(scored), 3) if scored else 0.0
    pending = [z for z in zones if z["status"] == "pending"]
    if not pending:
        complex_notes = []
    if not pending:
        status = "ready"
    elif len(pending) <= 2 and score >= 0.75 and not complex_notes:
        status = "needs_answer"
    else:
        status = "complex" if complex_notes or score < 0.75 else "needs_answer"
    return {
        "version": SPEC_VERSION,
        "kind": kind,
        "status": status,
        "recognition_score": score,
        "zones": zones,
        "sections": [{"heading": s["heading"], "type": s["type"], "title": s["title"]} for s in secs],
        "constraints": _constraints(doc, zones, kind),
        "complex_notes": sorted(set(complex_notes)),
        "fingerprint": fingerprint(doc, secs),
        "esn_name_guess": _esn_name(doc, ctx),
    }


def _template_kind(zones: list[dict]) -> str:
    samples = [z.get("sample") or "" for z in zones]
    ph = sum(placeholder_like(s) or F.is_placeholder(s) or bool(re.search(r"MM\s*/\s*(AAAA|YYYY)|XXXX", s)) for s in samples)
    return "blank" if ph >= 2 or (samples and ph / max(1, len(samples)) > 0.3) else "filled_example"


def _constraints(doc, zones: list[dict], kind: str = "filled_example") -> dict:
    font = None
    try:
        font = doc.styles["Normal"].font.name
    except KeyError:
        pass
    name_zone = next((z for z in zones if z["field"] == "nom_affiche"), None)
    mission = next((z for z in zones if z["field"] == "missions"), None)
    generic_clients = False
    if mission and mission.get("sample"):
        generic_clients = bool(GENERIC_CLIENT_WORDS.search(re.sub(r"\[[^\]]*\]", " ", mission["sample"])))
    sec = doc.sections[0]
    anonymize = bool(name_zone and name_zone["extra"].get("initials"))
    fields = {f for z in zones for f in str(z.get("field") or "").split(",")}
    if kind == "blank":
        generic_clients = False       # « CLIENT » d'un modèle vide est un emplacement, pas un client anonymisé
        anonymize = "trigramme" in fields and not ({"nom", "prenom", "nom_complet", "nom_affiche"} & fields)
    return {
        "font": font,
        "anonymize": anonymize,
        "anonymize_clients": generic_clients,
        "page": {"width_mm": round(sec.page_width.mm) if sec.page_width else None,
                 "height_mm": round(sec.page_height.mm) if sec.page_height else None},
        "max_pages": None,  # renseigné par le pipeline (pages du modèle d'origine)
    }


def fingerprint(doc, secs: list[dict]) -> str:
    """Empreinte structurelle : reconnaît le même modèle d'ESN quel que soit le consultant d'exemple."""
    parts = [f"{s['type'] or norm(s['title'])}" for s in secs]
    try:
        parts.append(doc.styles["Normal"].font.name or "")
    except KeyError:
        pass
    for s in doc.sections:
        for hp in s.header.paragraphs + s.footer.paragraphs:
            t = norm(hp.text)
            t = re.sub(r"\b[a-z]\.\s?[a-z]\.", "", t)
            parts.append(re.sub(r"\d+", "#", t))
        parts.append(f"{s.page_width}x{s.page_height}")
    return hashlib.sha256("|".join(parts).encode()).hexdigest()


def _esn_name(doc, ctx: _Ctx) -> str | None:
    for s in doc.sections:
        for hp in s.header.paragraphs:
            t = hp.text.strip()
            if t:
                return re.split(r"\s[—–-]\s|\|", t)[0].strip()[:80]
    for i in range(min(3, len(ctx.els))):
        t = ctx.text(i).strip()
        if t and t.isupper() and len(t) <= 40:
            return t.title()
    return None


def summary(spec: dict) -> Counter:
    return Counter((z["kind"], z["field"], z["status"]) for z in spec["zones"])
